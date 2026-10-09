"""Native text extraction for office and plain-text file formats.

Supported:
  .docx  → python-docx (paragraphs + tables), fallback to docx2txt
  .pptx  → python-pptx (all slide text frames)
  .txt   → raw bytes decoded as UTF-8 / latin-1

Not supported without system tools (LibreOffice / antiword):
  .doc   → raises DocumentLoadingError; convert to .docx first
  .ppt   → raises DocumentLoadingError; convert to .pptx first
"""

from __future__ import annotations

import io
import logging

logger = logging.getLogger(__name__)

# Public constant — imported by service.py and file_loader.py
TEXT_NATIVE_FILE_TYPES: frozenset[str] = frozenset({"docx", "doc", "pptx", "ppt", "txt"})


def extract_text_from_office_file(file_bytes: bytes, file_type: str) -> str:
    """Extract plain text from an office or plain-text document.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        file_type:  Lowercase extension without dot (e.g. ``"docx"``).

    Returns:
        Extracted text string (may be empty if the document has no readable text).

    Raises:
        DocumentLoadingError: If the file cannot be read or the format requires
            system-level tools that are not installed.
    """
    # Import here to avoid a circular import through app.config
    from parser.file_loader import DocumentLoadingError  # noqa: PLC0415

    dispatch = {
        "docx": _extract_docx,
        "pptx": _extract_pptx,
        "txt":  _extract_txt,
        "doc":  _unsupported_doc,
        "ppt":  _unsupported_ppt,
    }

    handler = dispatch.get(file_type)
    if handler is None:
        raise DocumentLoadingError(f"No text extractor registered for file type: {file_type!r}")

    return handler(file_bytes)


# ---------------------------------------------------------------------------
# Internal extractors
# ---------------------------------------------------------------------------

def _extract_docx(file_bytes: bytes) -> str:
    """Extract native text from a .docx file using python-docx with docx2txt fallback.

    OCR fallback for image-only documents is handled at a higher level by
    ``parser.enterprise_loader``.
    """
    from parser.file_loader import DocumentLoadingError

    try:
        import docx  # python-docx
        document = docx.Document(io.BytesIO(file_bytes))
        lines: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        lines.append(text)

        result = "\n".join(lines)
        if result:
            logger.info("Extracted %d characters from DOCX via python-docx", len(result))
            return result

    except Exception as primary_exc:
        logger.warning("python-docx extraction failed (%s); trying docx2txt", primary_exc)

    # Fallback: docx2txt
    try:
        import docx2txt
        result = docx2txt.process(io.BytesIO(file_bytes)) or ""
        logger.info("Extracted %d characters from DOCX via docx2txt", len(result))
        return result
    except Exception as fallback_exc:
        raise DocumentLoadingError("Unable to read DOCX file") from fallback_exc


def _extract_pptx(file_bytes: bytes) -> str:
    """Extract all text frames from a .pptx file using python-pptx."""
    from parser.file_loader import DocumentLoadingError

    try:
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(file_bytes))
        lines: list[str] = []

        for slide_index, slide in enumerate(presentation.slides, start=1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        lines.append(text)

        result = "\n".join(lines)
        logger.info(
            "Extracted %d characters from PPTX (%d slide(s))",
            len(result),
            len(presentation.slides),
        )
        return result
    except Exception as exc:
        raise DocumentLoadingError("Unable to read PPTX file") from exc


def _extract_txt(file_bytes: bytes) -> str:
    """Decode plain-text bytes trying common encodings."""
    from parser.file_loader import DocumentLoadingError

    for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            text = file_bytes.decode(encoding)
            logger.info("Decoded TXT as %s (%d characters)", encoding, len(text))
            return text
        except (UnicodeDecodeError, LookupError):
            continue
    raise DocumentLoadingError("Unable to decode text file: unsupported encoding")


def _unsupported_doc(_file_bytes: bytes) -> str:
    from parser.file_loader import DocumentLoadingError
    raise DocumentLoadingError(
        "Old binary .doc format is not supported. "
        "Please convert the file to .docx using Microsoft Word or LibreOffice first."
    )


def _unsupported_ppt(_file_bytes: bytes) -> str:
    from parser.file_loader import DocumentLoadingError
    raise DocumentLoadingError(
        "Old binary .ppt format is not supported. "
        "Please convert the file to .pptx using Microsoft PowerPoint or LibreOffice first."
    )
