"""Tests for parser/enterprise_loader.py — unified document text extraction."""

from __future__ import annotations

import io
import zipfile

import pytest

from parser.enterprise_loader import (
    MIN_NATIVE_CHARS,
    _extract_zip_images,
    _load_single_image,
    _normalise_file_type,
    load_document_text,
)
from utils.file_utils import UnsupportedFileTypeError, FileValidationError


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    import docx, io
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pptx_bytes(texts: list[str]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for text in texts:
        slide = prs.slides.add_slide(blank)
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        tb.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_zip_with_image(prefix: str, image_name: str) -> bytes:
    """Build a minimal ZIP archive with a 1×1 white PNG stored at prefix+image_name."""
    from PIL import Image as PILImage
    img_buf = io.BytesIO()
    PILImage.new("RGB", (1, 1), (255, 255, 255)).save(img_buf, format="PNG")
    img_buf.seek(0)
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w") as zf:
        zf.writestr(prefix + image_name, img_buf.read())
    return zip_buf.getvalue()


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

def test_min_native_chars_is_30() -> None:
    assert MIN_NATIVE_CHARS == 30


# ---------------------------------------------------------------------------
# _normalise_file_type
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("ext,expected", [
    (".pdf", "pdf"),
    (".docx", "docx"),
    (".doc", "doc"),
    (".pptx", "pptx"),
    (".ppt", "ppt"),
    (".txt", "txt"),
    (".jpg", "image"),
    (".jpeg", "image"),
    (".png", "image"),
    (".tif", "image"),
    (".tiff", "image"),
])
def test_normalise_file_type(ext: str, expected: str) -> None:
    assert _normalise_file_type(ext) == expected


# ---------------------------------------------------------------------------
# Unsupported / invalid inputs
# ---------------------------------------------------------------------------

def test_unsupported_extension_raises() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        load_document_text(b"data", "file.xlsx")


def test_file_too_large_raises() -> None:
    big_bytes = b"A" * (11 * 1024 * 1024)  # 11 MB > 10 MB limit
    with pytest.raises(FileValidationError):
        load_document_text(big_bytes, "big.txt")


# ---------------------------------------------------------------------------
# DOCX — sufficient native text skips OCR
# ---------------------------------------------------------------------------

def test_docx_sufficient_text_skips_ocr(monkeypatch) -> None:
    """When native text >= MIN_NATIVE_CHARS OCR must not be called."""
    from parser import enterprise_loader

    def _fail(_bytes, _type):
        raise AssertionError("OCR must not run when native text is sufficient")

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", _fail)
    long_text = "This document contains enough text to exceed the thirty-character threshold easily."
    text, file_type = load_document_text(_make_docx_bytes([long_text]), "doc.docx")
    assert long_text in text
    assert file_type == "docx"


# ---------------------------------------------------------------------------
# DOCX — short / empty native text triggers OCR
# ---------------------------------------------------------------------------

def test_docx_short_text_triggers_ocr(monkeypatch) -> None:
    """DOCX with < 30 chars of native text must invoke OCR fallback."""
    from parser import enterprise_loader

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", lambda _b, _t: "scanned diagram text")
    text, file_type = load_document_text(_make_docx_bytes(["Hi"]), "scan.docx")
    assert "scanned diagram text" in text
    assert file_type == "docx"


def test_docx_empty_text_uses_ocr_only(monkeypatch) -> None:
    """An empty DOCX returns only the OCR result."""
    from parser import enterprise_loader

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", lambda _b, _t: "ocr only result here")
    text, _ = load_document_text(_make_docx_bytes([]), "empty.docx")
    assert "ocr only result here" in text


def test_docx_short_text_prepended_to_ocr(monkeypatch) -> None:
    """Short native text is preserved and prepended before OCR output."""
    from parser import enterprise_loader

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", lambda _b, _t: "ocr appended text")
    text, _ = load_document_text(_make_docx_bytes(["Title"]), "diagram.docx")
    assert "Title" in text
    assert "ocr appended text" in text
    assert text.index("Title") < text.index("ocr appended text")


def test_docx_ocr_returns_empty_native_text_preserved(monkeypatch) -> None:
    """When OCR produces nothing, native text (even short) is returned as-is."""
    from parser import enterprise_loader

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", lambda _b, _t: "")
    text, _ = load_document_text(_make_docx_bytes(["Hi"]), "doc.docx")
    assert "Hi" in text


# ---------------------------------------------------------------------------
# PPTX — sufficient text skips OCR
# ---------------------------------------------------------------------------

def test_pptx_sufficient_text_skips_ocr(monkeypatch) -> None:
    from parser import enterprise_loader

    def _fail(_bytes, _type):
        raise AssertionError("OCR must not run for text-rich PPTX")

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", _fail)
    rich_text = "This presentation slide has more than thirty characters of content."
    text, file_type = load_document_text(_make_pptx_bytes([rich_text]), "deck.pptx")
    assert rich_text in text
    assert file_type == "pptx"


# ---------------------------------------------------------------------------
# TXT — sufficient text skips OCR
# ---------------------------------------------------------------------------

def test_txt_sufficient_text_skips_ocr(monkeypatch) -> None:
    from parser import enterprise_loader

    def _fail(_bytes, _type):
        raise AssertionError("OCR must not run for plain text file")

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", _fail)
    payload = "This plain text file has well over thirty characters of content.".encode("utf-8")
    text, file_type = load_document_text(payload, "notes.txt")
    assert "plain text file" in text
    assert file_type == "txt"


# ---------------------------------------------------------------------------
# Image — no native text, always goes to OCR
# ---------------------------------------------------------------------------

def test_image_has_no_native_text_triggers_ocr(monkeypatch) -> None:
    """Raw image files have no native text, so OCR always runs."""
    from parser import enterprise_loader
    from PIL import Image as PILImage

    # Build a real 1×1 PNG
    buf = io.BytesIO()
    PILImage.new("RGB", (1, 1), (0, 0, 0)).save(buf, format="PNG")
    png_bytes = buf.getvalue()

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", lambda _b, _t: "image ocr text")
    text, file_type = load_document_text(png_bytes, "photo.png")
    assert "image ocr text" in text
    assert file_type == "image"


# ---------------------------------------------------------------------------
# PDF — native text path (mocked)
# ---------------------------------------------------------------------------

def test_pdf_with_embedded_text_skips_ocr(monkeypatch) -> None:
    """PDF with enough embedded text must not trigger OCR."""
    from parser import enterprise_loader

    monkeypatch.setattr(
        enterprise_loader, "_extract_native_text",
        lambda _b, _t: "Invoice total amount due date payable by registered company name address",
    )

    def _fail(_bytes, _type):
        raise AssertionError("OCR must not run when PDF has sufficient embedded text")

    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", _fail)
    # %PDF bytes + validate_file_size passes; ensure_supported_extension needs .pdf
    text, file_type = load_document_text(b"%PDF-1.4 fake content", "doc.pdf")
    assert "Invoice" in text
    assert file_type == "pdf"


def test_pdf_image_only_triggers_ocr(monkeypatch) -> None:
    """A scanned PDF with no embedded text must fall back to OCR."""
    from parser import enterprise_loader

    monkeypatch.setattr(enterprise_loader, "_extract_native_text", lambda _b, _t: "")
    monkeypatch.setattr(enterprise_loader, "_extract_ocr_text", lambda _b, _t: "ocr from scanned pdf")
    text, file_type = load_document_text(b"%PDF-1.4 fake content", "scan.pdf")
    assert "ocr from scanned pdf" in text
    assert file_type == "pdf"


# ---------------------------------------------------------------------------
# _extract_zip_images — graceful handling of bad ZIP
# ---------------------------------------------------------------------------

def test_extract_zip_images_bad_zip_returns_empty() -> None:
    result = _extract_zip_images(b"this is not a zip", prefix="word/media/")
    assert result == []


def test_extract_zip_images_finds_images() -> None:
    zip_bytes = _make_zip_with_image("word/media/", "image1.png")
    images = _extract_zip_images(zip_bytes, prefix="word/media/")
    assert len(images) == 1


def test_extract_zip_images_ignores_non_image_files() -> None:
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w") as zf:
        zf.writestr("word/media/document.xml", "<xml/>")
    images = _extract_zip_images(zip_buf.getvalue(), prefix="word/media/")
    assert images == []


# ---------------------------------------------------------------------------
# _load_single_image — graceful handling of invalid bytes
# ---------------------------------------------------------------------------

def test_load_single_image_bad_bytes_returns_empty() -> None:
    result = _load_single_image(b"not an image")
    assert result == []


def test_load_single_image_valid_png() -> None:
    from PIL import Image as PILImage
    buf = io.BytesIO()
    PILImage.new("RGB", (4, 4), (128, 128, 128)).save(buf, format="PNG")
    result = _load_single_image(buf.getvalue())
    assert len(result) == 1
