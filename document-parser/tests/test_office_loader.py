"""Tests for parser/office_loader.py — native text extraction for DOCX/PPTX/TXT."""

from __future__ import annotations

import io

import pytest

from parser.file_loader import DocumentLoadingError
from parser.office_loader import TEXT_NATIVE_FILE_TYPES, extract_text_from_office_file


# ---------------------------------------------------------------------------
# Helpers: build minimal valid DOCX / PPTX in-memory
# ---------------------------------------------------------------------------

def _make_docx_bytes(paragraphs: list[str], table_cells: list[str] | None = None) -> bytes:
    """Build a minimal DOCX byte string using python-docx."""
    import docx
    document = docx.Document()
    for para in paragraphs:
        document.add_paragraph(para)
    if table_cells:
        table = document.add_table(rows=1, cols=len(table_cells))
        for i, text in enumerate(table_cells):
            table.cell(0, i).text = text
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_pptx_bytes(slide_texts: list[str]) -> bytes:
    """Build a minimal PPTX byte string using python-pptx."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # blank slide layout
    for text in slide_texts:
        slide = prs.slides.add_slide(blank_layout)
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        txBox.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# TEXT_NATIVE_FILE_TYPES constant
# ---------------------------------------------------------------------------

def test_text_native_types_contains_expected() -> None:
    assert TEXT_NATIVE_FILE_TYPES == {"docx", "doc", "pptx", "ppt", "txt"}


# ---------------------------------------------------------------------------
# DOCX extraction
# ---------------------------------------------------------------------------

def test_docx_extracts_paragraphs() -> None:
    file_bytes = _make_docx_bytes(["Hello World", "Second paragraph"])
    text = extract_text_from_office_file(file_bytes, "docx")
    assert "Hello World" in text
    assert "Second paragraph" in text


def test_docx_extracts_table_cells() -> None:
    file_bytes = _make_docx_bytes(["Header"], table_cells=["Cell A", "Cell B"])
    text = extract_text_from_office_file(file_bytes, "docx")
    assert "Cell A" in text
    assert "Cell B" in text


def test_docx_empty_document_returns_string() -> None:
    file_bytes = _make_docx_bytes([])
    result = extract_text_from_office_file(file_bytes, "docx")
    assert isinstance(result, str)


def test_docx_invalid_bytes_raises_document_loading_error() -> None:
    with pytest.raises(DocumentLoadingError, match="DOCX"):
        extract_text_from_office_file(b"not a docx file at all", "docx")


# ---------------------------------------------------------------------------
# PPTX extraction
# ---------------------------------------------------------------------------

def test_pptx_extracts_slide_text() -> None:
    file_bytes = _make_pptx_bytes(["Slide one content", "Slide two content"])
    text = extract_text_from_office_file(file_bytes, "pptx")
    assert "Slide one content" in text
    assert "Slide two content" in text


def test_pptx_empty_presentation_returns_string() -> None:
    file_bytes = _make_pptx_bytes([])
    result = extract_text_from_office_file(file_bytes, "pptx")
    assert isinstance(result, str)


def test_pptx_invalid_bytes_raises_document_loading_error() -> None:
    with pytest.raises(DocumentLoadingError, match="PPTX"):
        extract_text_from_office_file(b"not a pptx", "pptx")


# ---------------------------------------------------------------------------
# TXT extraction
# ---------------------------------------------------------------------------

def test_txt_utf8_extraction() -> None:
    text = extract_text_from_office_file("Hello UTF-8 text".encode("utf-8"), "txt")
    assert "Hello UTF-8 text" in text


def test_txt_latin1_extraction() -> None:
    text = extract_text_from_office_file("Caf\xe9 document".encode("latin-1"), "txt")
    assert "Caf" in text


def test_txt_utf8_bom_extraction() -> None:
    text = extract_text_from_office_file("\ufeffBOM prefixed text".encode("utf-8-sig"), "txt")
    assert "BOM prefixed text" in text


# ---------------------------------------------------------------------------
# DOC / PPT — unsupported
# ---------------------------------------------------------------------------

def test_doc_raises_document_loading_error_with_guidance() -> None:
    with pytest.raises(DocumentLoadingError, match=r"\.docx"):
        extract_text_from_office_file(b"\xd0\xcf\x11\xe0", "doc")  # typical binary .doc header


def test_ppt_raises_document_loading_error_with_guidance() -> None:
    with pytest.raises(DocumentLoadingError, match=r"\.pptx"):
        extract_text_from_office_file(b"\xd0\xcf\x11\xe0", "ppt")


# ---------------------------------------------------------------------------
# Unknown type
# ---------------------------------------------------------------------------

def test_unknown_file_type_raises_document_loading_error() -> None:
    with pytest.raises(DocumentLoadingError, match="No text extractor"):
        extract_text_from_office_file(b"data", "xlsx")


# ---------------------------------------------------------------------------
# file_loader integration — new extensions routed correctly
# ---------------------------------------------------------------------------

def test_file_loader_returns_empty_images_for_docx() -> None:
    from parser.file_loader import load_document_as_images
    file_bytes = _make_docx_bytes(["Test content for routing"])
    images, file_type = load_document_as_images(file_bytes, "document.docx")
    assert images == []
    assert file_type == "docx"


def test_file_loader_returns_empty_images_for_pptx() -> None:
    from parser.file_loader import load_document_as_images
    file_bytes = _make_pptx_bytes(["Slide text"])
    images, file_type = load_document_as_images(file_bytes, "presentation.pptx")
    assert images == []
    assert file_type == "pptx"


def test_file_loader_returns_empty_images_for_txt() -> None:
    from parser.file_loader import load_document_as_images
    images, file_type = load_document_as_images(b"plain text content", "notes.txt")
    assert images == []
    assert file_type == "txt"


def test_file_loader_still_rejects_unsupported_extension() -> None:
    from parser.file_loader import load_document_as_images
    from utils.file_utils import UnsupportedFileTypeError
    with pytest.raises(UnsupportedFileTypeError):
        load_document_as_images(b"data", "file.xlsx")


# ---------------------------------------------------------------------------
# config — new extensions present
# ---------------------------------------------------------------------------

def test_config_includes_new_extensions() -> None:
    from app.config import settings
    for ext in (".docx", ".doc", ".pptx", ".ppt", ".txt"):
        assert ext in settings.supported_extensions, f"{ext} missing from supported_extensions"
